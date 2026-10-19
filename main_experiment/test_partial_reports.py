"""Partial validation is usable feedback, never complete-delivery evidence."""
from pathlib import Path
import tempfile
import unittest
import openpyxl

from main_experiment.evaluate import conformance_metrics

PINS=['FDA 2508.1','1.0.9','2026-03-27','ADaM-IG 1.3','2.1','MedDRA','8.0']


def workbook(path,warning=False):
    book=openpyxl.Workbook()
    summary=book.active;summary.title='Validation Summary'
    summary.append([' '.join(PINS)])
    if warning:
        summary.append(['Problems with your Validator installation detected'])
        summary.append(['Problem'])
        summary.append(['MedDRA not configured'])
    datasets=book.create_sheet('Dataset Summary')
    datasets.append(['Domain','Label','Class','Source','Records','Rejects'])
    datasets.append(['ADSL','Subjects','ADAM','adsl.xpt',2,0])
    issues=book.create_sheet('Issue Summary')
    issues.append(['Source','Pinnacle 21 ID','Message','Severity','Found'])
    issues.append(['ADSL','AD0001','Synthetic finding','Warning',2])
    book.save(path)


class PartialReportTests(unittest.TestCase):
    def test_partial_report_usable_without_complete_deliverables(self):
        with tempfile.TemporaryDirectory() as t:
            p=Path(t)/'report.xlsx';workbook(p)
            result=conformance_metrics(p,['ADSL','ADAE'],['ADSL'],PINS)
            self.assertTrue(result['report_usable'])
            self.assertEqual('valid_partial',result['status'])
            self.assertFalse(result['validation_coverage_complete'])
            self.assertEqual(['ADAE'],result['missing_datasets'])
            self.assertEqual(['ADSL'],result['validated_datasets'])
            self.assertEqual(1,result['dataset_rule_pairs'])

    def test_subset_of_produced_domains_is_still_usable(self):
        with tempfile.TemporaryDirectory() as t:
            p=Path(t)/'report.xlsx';workbook(p)
            result=conformance_metrics(p,['ADSL','ADAE'],['ADSL','ADAE'],PINS)
            self.assertTrue(result['report_usable'])
            self.assertFalse(result['validation_coverage_complete'])
            self.assertEqual(['ADAE'],result['produced_not_validated'])
            self.assertEqual('valid_partial',result['status'])

    def test_full_report_remains_valid(self):
        with tempfile.TemporaryDirectory() as t:
            p=Path(t)/'report.xlsx';workbook(p)
            result=conformance_metrics(p,['ADSL'],['ADSL'],PINS)
            self.assertTrue(result['report_usable'])
            self.assertTrue(result['validation_coverage_complete'])
            self.assertEqual('valid',result['status'])

    def test_installation_version_and_observed_coverage_still_required(self):
        with tempfile.TemporaryDirectory() as t:
            p=Path(t)/'report.xlsx'
            for warning,pins,actual in [(True,PINS,['ADSL']),
                                       (False,PINS+['WRONG-VERSION'],['ADSL']),
                                       (False,PINS,['ADAE']),
                                       (False,PINS,[])]:
                workbook(p,warning)
                result=conformance_metrics(p,['ADSL','ADAE'],actual,pins)
                self.assertFalse(result['report_usable'])
                self.assertIsNone(result['dataset_rule_pairs'])
                self.assertEqual(1,result['observed_dataset_rule_pairs'])

    def test_corrupt_report_is_not_usable(self):
        with tempfile.TemporaryDirectory() as t:
            p=Path(t)/'report.xlsx';p.write_bytes(b'not an xlsx')
            with self.assertRaises(Exception):
                conformance_metrics(p,['ADSL'],['ADSL'],PINS)
